# generate_report.py
"""
Master Orchestration Script for Generating the Complete Project Report:
1. Builds full editable Microsoft Word document (.docx)
   following university academic engineering report standards.
2. Uses Microsoft Word COM automation to update fields and export
   high-fidelity PDF document (.pdf).
"""

import os
import sys
import time
import docx
import win32com.client
from report_builder.styles import setup_document_styles
from report_builder.front_matter import build_front_matter
from report_builder.chapter1 import build_chapter1
from report_builder.chapter2 import build_chapter2
from report_builder.chapter3 import build_chapter3
from report_builder.chapter4 import build_chapter4
from report_builder.chapter5 import build_chapter5
from report_builder.chapter6 import build_chapter6
from report_builder.chapter7 import build_chapter7
from report_builder.chapter8 import build_chapter8
from report_builder.chapter9 import build_chapter9
from report_builder.chapter10 import build_chapter10
from report_builder.references_and_appendices import build_references_and_appendices

OUTPUT_DIR = r"e:\Data Science Project"
DOCX_NAME = "Retail_Footfall_Queue_Monitoring_Project_Report.docx"
PDF_NAME = "Retail_Footfall_Queue_Monitoring_Project_Report.pdf"

DOCX_PATH = os.path.join(OUTPUT_DIR, DOCX_NAME)
PDF_PATH = os.path.join(OUTPUT_DIR, PDF_NAME)


def generate_word_document():
    """Assembles all report sections into the Word document."""
    print("=" * 70)
    print("Starting generation of Retail Footfall Queue Report...")
    print("=" * 70)

    doc = docx.Document()
    setup_document_styles(doc)

    print("[1/12] Building Front Matter (Cover, Certificate, TOC, Lists)...")
    build_front_matter(doc)

    print("[2/12] Building Chapter 1: Introduction...")
    build_chapter1(doc)

    print("[3/12] Building Chapter 2: Literature Survey...")
    build_chapter2(doc)

    print("[4/12] Building Chapter 3: Existing vs. Proposed System...")
    build_chapter3(doc)

    print("[5/12] Building Chapter 4: System Architecture and Design...")
    build_chapter4(doc)

    print("[6/12] Building Chapter 5: Methodology and Mathematical Formulation...")
    build_chapter5(doc)

    print("[7/12] Building Chapter 6: Machine Learning for Queue Forecasting...")
    build_chapter6(doc)

    print("[8/12] Building Chapter 7: Dataset Description and Exploratory Data Analysis...")
    build_chapter7(doc)

    print("[9/12] Building Chapter 8: System Implementation and Dashboard Walkthrough...")
    build_chapter8(doc)

    print("[10/12] Building Chapter 9: Experimental Results and Performance Evaluation...")
    build_chapter9(doc)

    print("[11/12] Building Chapter 10: Conclusion and Future Enhancements...")
    build_chapter10(doc)

    print("[12/12] Building References and Appendices A & B...")
    build_references_and_appendices(doc)

    print(f"Saving Word document to: {DOCX_PATH}...")
    doc.save(DOCX_PATH)
    docx_size = os.path.getsize(DOCX_PATH) / 1024
    print(f"Successfully created {DOCX_NAME} ({docx_size:.1f} KB)!")

    # Count words and paragraphs
    total_words = sum(len(p.text.split()) for p in doc.paragraphs)
    total_paragraphs = len(doc.paragraphs)
    total_tables = len(doc.tables)
    print(f"Document Metrics: {total_words:,} words | {total_paragraphs} paragraphs | {total_tables} tables")


def export_to_pdf_via_word_com():
    """Automates Microsoft Word to update fields and export PDF."""
    print("=" * 70)
    print("Exporting Word document to PDF via Microsoft Word COM automation...")
    print("=" * 70)

    if not os.path.exists(DOCX_PATH):
        raise FileNotFoundError(f"Source DOCX file not found: {DOCX_PATH}")

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    word.DisplayAlerts = False

    try:
        abs_docx = os.path.abspath(DOCX_PATH)
        abs_pdf = os.path.abspath(PDF_PATH)

        print(f"Opening Word document: {abs_docx}...")
        wb = word.Documents.Open(abs_docx)

        print("Updating document fields (TOC, page numbers, cross-references)...")
        try:
            wb.Fields.Update()
        except Exception as e:
            print(f"Fields update note: {e}")

        # Save the updated docx with refreshed fields
        wb.Save()

        # 17 = wdFormatPDF
        print(f"Exporting PDF to: {abs_pdf}...")
        wb.SaveAs(abs_pdf, FileFormat=17)

        page_count = wb.ComputeStatistics(2)  # 2 = wdStatisticPages
        wb.Close()

        pdf_size = os.path.getsize(PDF_PATH) / (1024 * 1024)
        print(f"Successfully generated PDF: {PDF_NAME}")
        print(f"PDF Metrics: {page_count} pages | {pdf_size:.2f} MB")

    finally:
        word.Quit()


if __name__ == "__main__":
    generate_word_document()
    export_to_pdf_via_word_com()
    print("=" * 70)
    print("ALL DELIVERABLES SUCCESSFULLY GENERATED!")
    print(f"1. Microsoft Word Document: {DOCX_PATH}")
    print(f"2. PDF Document:            {PDF_PATH}")
    print("=" * 70)
