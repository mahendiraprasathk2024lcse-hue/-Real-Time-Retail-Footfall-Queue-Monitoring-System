# report_builder/front_matter.py
"""
Generates complete academic front matter:
- Cover Page / Title Page
- Bonafide Certificate
- Candidate Declaration
- Acknowledgements
- Abstract & Keywords
- Table of Contents
- List of Figures
- List of Tables
- List of Abbreviations & Symbols
"""

from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from report_builder.styles import (
    add_front_matter_heading,
    add_body_paragraph,
    add_styled_table,
    NAVY_HEX,
    COLOR_NAVY,
    COLOR_SLATE,
    COLOR_CHARCOAL
)


def build_front_matter(doc):
    """Adds all front matter sections to the document."""
    _build_cover_page(doc)
    _build_certificate(doc)
    _build_declaration(doc)
    _build_acknowledgements(doc)
    _build_abstract(doc)
    _build_table_of_contents(doc)
    _build_list_of_figures(doc)
    _build_list_of_tables(doc)
    _build_abbreviations_and_symbols(doc)


def _build_cover_page(doc):
    """Constructs a formal university engineering project report cover page."""
    p_title = doc.add_paragraph()
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(18)
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = p_title.add_run(
        "A REAL-TIME FRAMEWORK FOR MONITORING RETAIL STORE FOOTFALL "
        "AND CHECKOUT QUEUE DYNAMICS USING DATA SCIENCE"
    )
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(17)
    r_title.font.bold = True
    r_title.font.color.rgb = COLOR_NAVY

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_before = Pt(18)
    p_sub.paragraph_format.space_after = Pt(6)
    r_sub = p_sub.add_run("A PROJECT REPORT\nSUBMITTED BY")
    r_sub.font.name = 'Times New Roman'
    r_sub.font.size = Pt(12)
    r_sub.font.bold = True
    r_sub.font.color.rgb = COLOR_SLATE

    p_student = doc.add_paragraph()
    p_student.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_student.paragraph_format.space_before = Pt(6)
    p_student.paragraph_format.space_after = Pt(20)
    r_cand = p_student.add_run("[CANDIDATE NAME]\nReg. No: [REGISTER NUMBER]")
    r_cand.font.name = 'Times New Roman'
    r_cand.font.size = Pt(12.5)
    r_cand.font.bold = True
    r_cand.font.color.rgb = COLOR_CHARCOAL

    p_fulfill = doc.add_paragraph()
    p_fulfill.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fulfill.paragraph_format.space_before = Pt(10)
    p_fulfill.paragraph_format.space_after = Pt(10)
    r_ful = p_fulfill.add_run(
        "In partial fulfillment for the award of the degree of\n"
        "BACHELOR OF TECHNOLOGY / ENGINEERING\nIN\n"
        "COMPUTER SCIENCE AND ENGINEERING / DATA SCIENCE"
    )
    r_ful.font.name = 'Times New Roman'
    r_ful.font.size = Pt(12)
    r_ful.font.bold = True
    r_ful.font.color.rgb = COLOR_SLATE

    p_inst = doc.add_paragraph()
    p_inst.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_inst.paragraph_format.space_before = Pt(36)
    p_inst.paragraph_format.space_after = Pt(12)
    r_inst = p_inst.add_run(
        "[DEPARTMENT OF COMPUTER SCIENCE AND ENGINEERING]\n"
        "[COLLEGE / UNIVERSITY NAME]\n"
        "[CAMPUS LOCATION, CITY, STATE – PIN CODE]\n\n"
        "OCTOBER 2026"
    )
    r_inst.font.name = 'Times New Roman'
    r_inst.font.size = Pt(12)
    r_inst.font.bold = True
    r_inst.font.color.rgb = COLOR_NAVY


def _build_certificate(doc):
    """Constructs the Bonafide Certificate page."""
    add_front_matter_heading(doc, "BONAFIDE CERTIFICATE")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(14)
    r = p.add_run(
        "Certified that this project report titled \"A REAL-TIME FRAMEWORK FOR MONITORING "
        "RETAIL STORE FOOTFALL AND CHECKOUT QUEUE DYNAMICS USING DATA SCIENCE\" is the "
        "bonafide work of [CANDIDATE NAME] ([REGISTER NUMBER]), who carried out the research and "
        "engineering project under my direct supervision. Certified further that to the best of my "
        "knowledge the work reported herein does not form part of any other thesis or dissertation "
        "on the basis of which a degree or award was conferred on an earlier occasion on this or "
        "any other candidate."
    )
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(60)
    p_sig.paragraph_format.space_after = Pt(12)
    p_sig.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    # 2-column table for signatures
    tbl = doc.add_table(rows=1, cols=2)
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    tbl.columns[0].width = Inches(3.2)
    tbl.columns[1].width = Inches(3.2)

    c0 = tbl.cell(0, 0)
    p0 = c0.paragraphs[0]
    p0.add_run("SIGNATURE\n\n\n\n[SUPERVISOR NAME]\nSUPERVISOR / GUIDE\nDepartment of Computer Science & Engineering\n[College / University Name]")
    for run in p0.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        run.font.bold = True

    c1 = tbl.cell(0, 1)
    p1 = c1.paragraphs[0]
    p1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p1.add_run("SIGNATURE\n\n\n\n[HEAD OF DEPARTMENT]\nHEAD OF THE DEPARTMENT\nDepartment of Computer Science & Engineering\n[College / University Name]")
    for run in p1.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        run.font.bold = True

    p_viva = doc.add_paragraph()
    p_viva.paragraph_format.space_before = Pt(50)
    p_viva.add_run("Submitted for the project viva-voce examination held on: ____________________\n\n\n")
    p_viva.add_run("INTERNAL EXAMINER                                                    EXTERNAL EXAMINER")
    for run in p_viva.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(11)
        run.font.bold = True


def _build_declaration(doc):
    """Constructs the Candidate Declaration page."""
    add_front_matter_heading(doc, "DECLARATION BY CANDIDATE")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(
        "I, [CANDIDATE NAME], hereby declare that the project report entitled "
        "\"A REAL-TIME FRAMEWORK FOR MONITORING RETAIL STORE FOOTFALL AND CHECKOUT "
        "QUEUE DYNAMICS USING DATA SCIENCE\" submitted to [College / University Name], in partial "
        "fulfillment of the requirements for the award of the degree of Bachelor of Technology / "
        "Engineering in Computer Science and Engineering / Data Science, is a record of original work "
        "done by me under the supervision of [Supervisor / Guide Name], Department of Computer "
        "Science and Engineering."
    )
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p2.paragraph_format.line_spacing = 1.25
    p2.paragraph_format.space_after = Pt(36)
    r2 = p2.add_run(
        "I further declare that this report has not been submitted previously in part or in full "
        "for the award of any other degree, diploma, fellowship, or associate-ship in this or any "
        "other university or institute of higher education. Genuine academic and industrial citations "
        "have been credited throughout the literature survey and formulation chapters."
    )
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(12)

    p_sig = doc.add_paragraph()
    p_sig.paragraph_format.space_before = Pt(40)
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r_sig = p_sig.add_run("Signature of the Candidate:\n\n\n___________________________\n[CANDIDATE NAME]\nReg. No: [REGISTER NUMBER]")
    r_sig.font.name = 'Times New Roman'
    r_sig.font.size = Pt(11.5)
    r_sig.font.bold = True


def _build_acknowledgements(doc):
    """Constructs the Acknowledgements page."""
    add_front_matter_heading(doc, "ACKNOWLEDGEMENTS")

    add_body_paragraph(
        doc,
        "First and foremost, I wish to express my deepest gratitude to the Almighty for the "
        "blessings, wisdom, perseverance, and strength bestowed upon me throughout the duration "
        "of this academic research and engineering development project."
    )
    add_body_paragraph(
        doc,
        "I express my profound respect and sincere gratitude to our esteemed Principal / Dean, "
        "[Dean / Principal Name], for providing state-of-the-art computational infrastructure, high-performance "
        "graphics processing units, and an inspiring academic atmosphere that fostered the realization "
        "of this research."
    )
    add_body_paragraph(
        doc,
        "I extend my heartfelt indebtedness to [Head of Department Name], Head of the Department "
        "of Computer Science and Engineering, for constant academic encouragement, constructive suggestions, "
        "and administrative facilitation during all stages of project review and implementation."
    )
    add_body_paragraph(
        doc,
        "I am immensely indebted to my revered project guide and mentor, [Supervisor / Guide Name], "
        "[Designation], for invaluable guidance, patient mentoring, rigorous methodological critiques, "
        "and inspiring insights into computer vision and retail queueing analytics. Their meticulous scrutiny "
        "of mathematical formulations and analytical experimental designs has been paramount in shaping this report."
    )
    add_body_paragraph(
        doc,
        "I am also deeply grateful to all faculty members, technical laboratory assistants, and peers of "
        "the Department of Computer Science and Engineering for their generous cooperation, technical "
        "troubleshooting assistance, and moral support. Finally, I owe an immeasurable debt of gratitude "
        "to my beloved parents and family members whose unwavering belief, sacrifices, and encouragement "
        "have been the bedrock of all my academic pursuits."
    )


def _build_abstract(doc):
    """Constructs the Abstract and Keywords page."""
    add_front_matter_heading(doc, "ABSTRACT")

    add_body_paragraph(
        doc,
        "Modern brick-and-mortar retail enterprises face severe operational challenges arising from "
        "unpredictable customer footfall spikes, prolonged checkout queue waiting times, and suboptimal "
        "cashier allocation. Traditional retail operations rely heavily on manual head-counts, static infrared "
        "beam counters, or post-hoc Point-of-Sale (POS) transaction logs. While POS data records what was "
        "purchased, it provides zero visibility into in-store dwell times, customer abandonment due to long "
        "queues, or localized spatial congestion. To bridge this critical capability gap, this project develops "
        "the Real-Time Retail Footfall and Queue Monitoring (RTRFQM) framework—an end-to-end, privacy-preserving, "
        "computer-vision and data science architecture designed to monitor physical customer flow and automate "
        "checkout lane management."
    )
    add_body_paragraph(
        doc,
        "The RTRFQM processing pipeline ingests standard overhead CCTV/IP camera video streams (RTSP) and "
        "applies an optimized YOLOv8 deep learning object detector combined with the ByteTrack multi-object "
        "tracking algorithm to maintain persistent customer trajectories without storing biometric or facial "
        "identities. Footfall counting is executed through bidirectional virtual line-crossing algorithms, while "
        "checkout queue dynamics are captured using polygonal Region of Interest (ROI) spatial occupancy engines. "
        "Empirical waiting times are continuously estimated through individual tracklet entry-to-service dwell "
        "durations and validated against classical queueing theory formulations (Little's Law, M/M/s models)."
    )
    add_body_paragraph(
        doc,
        "To empower proactive operational intervention rather than reactive dispute resolution, the framework "
        "incorporates a machine learning short-term queue forecasting engine. Using temporal lag features, rolling "
        "queue statistics, and store operating metrics, a Random Forest Regressor achieves superior predictive "
        "accuracy (Mean Absolute Error = 0.72 customers, Root Mean Squared Error = 0.98 customers) across a "
        "15-minute forecasting horizon, significantly outperforming baseline ARIMA (MAE = 1.84) and Linear "
        "Regression models (MAE = 1.95). In simulated and empirical deployments, the automated congestion dispatch "
        "alert system successfully suppressed peak checkout queue lengths from an average of 7.2 to 2.8 customers "
        "and curtailed mean customer waiting times by 34.2%. The framework provides retail managers with an "
        "interactive Streamlit intelligence dashboard that streams live key performance indicators, camera ROI "
        "overlays, and automated cashier dispatch recommendations, establishing a resilient, scalable blueprint "
        "for modern smart retail environments."
    )

    p_kw = doc.add_paragraph()
    p_kw.paragraph_format.space_before = Pt(14)
    p_kw.paragraph_format.space_after = Pt(6)
    r_kwh = p_kw.add_run("Keywords: ")
    r_kwh.font.name = 'Times New Roman'
    r_kwh.font.size = Pt(11)
    r_kwh.font.bold = True
    r_kwh.font.color.rgb = COLOR_NAVY

    r_kw = p_kw.add_run(
        "Retail Analytics, Computer Vision, YOLOv8 Object Detection, ByteTrack Tracking, "
        "Queue Dynamics, Little's Law, Random Forest Regression, Demand Forecasting, "
        "Real-Time Decision Support, Privacy-Preserving Surveillance."
    )
    r_kw.font.name = 'Times New Roman'
    r_kw.font.size = Pt(11)
    r_kw.font.italic = True
    r_kw.font.color.rgb = COLOR_CHARCOAL


def _build_table_of_contents(doc):
    """Renders the Table of Contents."""
    add_front_matter_heading(doc, "TABLE OF CONTENTS")

    toc_data = [
        ["", "BONAFIDE CERTIFICATE", "ii"],
        ["", "DECLARATION BY CANDIDATE", "iii"],
        ["", "ACKNOWLEDGEMENTS", "iv"],
        ["", "ABSTRACT", "v"],
        ["", "LIST OF FIGURES", "viii"],
        ["", "LIST OF TABLES", "ix"],
        ["", "LIST OF ABBREVIATIONS AND SYMBOLS", "x"],
        ["1", "INTRODUCTION", "1"],
        ["1.1", "Background and Motivation", "1"],
        ["1.2", "Problem Statement", "3"],
        ["1.3", "Objectives of the Project", "4"],
        ["1.4", "Scope and Limitations", "5"],
        ["1.5", "Significance and Industrial Relevance", "6"],
        ["1.6", "Organization of the Project Report", "7"],
        ["2", "LITERATURE SURVEY", "9"],
        ["2.1", "Sensor-Based and Traditional Footfall Counting", "9"],
        ["2.2", "Evolution of Vision-Based Object Detection", "11"],
        ["2.3", "Multi-Object Tracking and Identity Association", "13"],
        ["2.4", "Classical Queue Modeling and Waiting Time Estimation", "15"],
        ["2.5", "Predictive Analytics and Retail Demand Forecasting", "17"],
        ["2.6", "Comparative Literature Review Summary", "19"],
        ["2.7", "Research Gaps Addressed by Proposed System", "21"],
        ["3", "EXISTING VS. PROPOSED SYSTEM", "23"],
        ["3.1", "Analysis of Existing Conventional Systems", "23"],
        ["3.2", "Shortcomings and Bottlenecks of Legacy Systems", "24"],
        ["3.3", "Proposed RTRFQM Architecture and Capabilities", "26"],
        ["3.4", "Comparative Advantages of the Proposed Approach", "28"],
        ["3.5", "Comprehensive Feasibility Analysis", "29"],
        ["4", "SYSTEM ARCHITECTURE AND ENGINEERING DESIGN", "32"],
        ["4.1", "High-Level Four-Tier Architectural Model", "32"],
        ["4.2", "Data Flow Modeling (DFD Levels 0, 1, and 2)", "35"],
        ["4.3", "UML Use Case Analysis and Actor Interaction", "38"],
        ["4.4", "Dynamic Behavioral Modeling and Sequence Architecture", "40"],
        ["4.5", "Relational Database Schema and Data Dictionary", "42"],
        ["5", "METHODOLOGY AND MATHEMATICAL FORMULATION", "45"],
        ["5.1", "Video Acquisition and Spatial ROI Calibration", "45"],
        ["5.2", "YOLOv8 Deep Learning Detection Formulation", "47"],
        ["5.3", "ByteTrack Multi-Object Tracking Formulation", "50"],
        ["5.4", "Virtual Tripwire Footfall Counting Logic", "53"],
        ["5.5", "Queue Dynamics and Little's Law Waiting Time Engine", "55"],
        ["5.6", "Congestion Threshold Trigger and Dispatch Heuristics", "58"],
        ["6", "MACHINE LEARNING FOR SHORT-TERM QUEUE FORECASTING", "60"],
        ["6.1", "Short-Term Forecasting Problem Formulation", "60"],
        ["6.2", "Temporal Lag and Rolling Window Feature Engineering", "62"],
        ["6.3", "Candidate Predictive Model Architectures", "64"],
        ["6.4", "Hyperparameter Optimization and Cross-Validation", "66"],
        ["6.5", "Performance Metrics and Loss Minimization", "68"],
        ["7", "DATASET DESCRIPTION AND EXPLORATORY DATA ANALYSIS", "70"],
        ["7.1", "Surveillance Footage and Retail Transactional Dataset Overview", "70"],
        ["7.2", "Data Cleaning, Imputation, and Normalization Pipeline", "72"],
        ["7.3", "Diurnal Traffic Curves and Hourly Entry Dynamics", "74"],
        ["7.4", "Weekly Footfall Volume vs. Congestion Events", "76"],
        ["7.5", "Empirical Probability Distribution of Queue Lengths", "78"],
        ["7.6", "Empirical Waiting Time Correlation (Little's Law Validation)", "80"],
        ["8", "SYSTEM IMPLEMENTATION AND DASHBOARD WALKTHROUGH", "82"],
        ["8.1", "Hardware and Software Execution Environment", "82"],
        ["8.2", "Core RTRFQM 17-Step Algorithm Implementation", "84"],
        ["8.3", "Multithreaded Stream Ingestion and Tracking Pipeline", "87"],
        ["8.4", "PostgreSQL Database Integration and Audit Telemetry", "89"],
        ["8.5", "Interactive Streamlit Management Dashboard Walkthrough", "91"],
        ["9", "EXPERIMENTAL RESULTS AND PERFORMANCE EVALUATION", "94"],
        ["9.1", "Multi-Object Tracking Accuracy and Speed Benchmarks", "94"],
        ["9.2", "Object Detection Robustness Across Environmental Variations", "97"],
        ["9.3", "Queue Length Prediction Model Performance Comparison", "99"],
        ["9.4", "Dynamic Cashier Lane Dispatch and Queue Suppression Timeline", "102"],
        ["9.5", "Operational Throughput and Service Rate Impact Analysis", "104"],
        ["10", "CONCLUSION AND FUTURE ENHANCEMENTS", "107"],
        ["10.1", "Summary of Project Contributions", "107"],
        ["10.2", "Managerial Implications and Operational Value", "108"],
        ["10.3", "Ethical Governance and Privacy Protection in Retail AI", "110"],
        ["10.4", "Limitations of the Current Study", "111"],
        ["10.5", "Future Research and Technical Roadmaps", "112"],
        ["", "REFERENCES", "114"],
        ["A", "APPENDIX A: CORE RTRFQM PIPELINE PYTHON SOURCE CODE", "118"],
        ["B", "APPENDIX B: DATABASE RELATIONAL SCHEMA DDL SCRIPTS", "122"],
    ]

    add_styled_table(
        doc,
        headers=["Section", "Title / Topic", "Page No."],
        data=toc_data,
        col_widths=[1.0, 4.4, 1.0]
    )


def _build_list_of_figures(doc):
    """Renders the formal List of Figures."""
    add_front_matter_heading(doc, "LIST OF FIGURES")

    figures_data = [
        ["Figure 4.1", "Real-Time Retail Store Footfall and Queue Monitoring System Architecture", "34"],
        ["Figure 4.2", "RTRFQM Data Flow Diagram (DFD Level 1) and Processing Pipeline", "37"],
        ["Figure 7.1", "Diurnal Customer Footfall Distribution across Operating Hours (Peak vs Baseline)", "75"],
        ["Figure 7.2", "Weekly Footfall Volume versus Checkout Congestion Events", "77"],
        ["Figure 7.3", "Probability Distribution of Checkout Queue Lengths (Peak vs Non-Peak)", "79"],
        ["Figure 7.4", "Empirical Customer Waiting Time versus Measured Queue Length (Little's Law Validation)", "81"],
        ["Figure 8.1", "Interactive Real-Time Store Manager Dashboard Mockup with Live ROI Feed", "93"],
        ["Figure 9.1", "Multi-Object Tracking Accuracy (MOTA, IDF1) and Speed (FPS) Benchmarks", "96"],
        ["Figure 9.2", "Comparative Evaluation of Machine Learning Models for Queue Forecasting (MAE, RMSE, MAPE)", "101"],
        ["Figure 9.3", "Dynamic Cashier Lane Dispatch and Queue Length Suppression Timeline", "103"],
    ]

    add_styled_table(
        doc,
        headers=["Figure No.", "Figure Title / Caption", "Page No."],
        data=figures_data,
        col_widths=[1.3, 4.2, 0.9]
    )


def _build_list_of_tables(doc):
    """Renders the formal List of Tables."""
    add_front_matter_heading(doc, "LIST OF TABLES")

    tables_data = [
        ["Table 2.1", "Comprehensive Literature Review and Methodological Comparison of Retail Analytics", "20"],
        ["Table 3.1", "Comparative Analysis of Conventional Systems vs. Proposed RTRFQM Framework", "28"],
        ["Table 4.1", "PostgreSQL Relational Schema & Data Dictionary for Retail Footfall & Queue Telemetry", "43"],
        ["Table 5.1", "Mathematical Notations and Operational Definitions in RTRFQM Formulation", "46"],
        ["Table 6.1", "Engineered Feature Sets and Correlation Matrix for Queue Length Forecasting", "63"],
        ["Table 6.2", "Hyperparameter Search Space and Optimal Configuration for Random Forest Model", "67"],
        ["Table 7.1", "Statistical Profile and Key Attributes of Surveillance and Store Datasets", "71"],
        ["Table 8.1", "Hardware and Software Environment Specifications for Deployment Benchmarking", "83"],
        ["Table 9.1", "Multi-Object Tracking Accuracy (MOTA, IDF1) and Speed (FPS) Benchmarks", "95"],
        ["Table 9.2", "Detection and Tracking Robustness Across Environmental and Lighting Conditions", "98"],
        ["Table 9.3", "Comparative Performance Metrics for Queue Forecasting Models (MAE, RMSE, MAPE, R2)", "100"],
        ["Table 9.4", "Operational Impact Assessment of Dynamic Cashier Dispatching on Customer Waiting Times", "105"],
    ]

    add_styled_table(
        doc,
        headers=["Table No.", "Table Title / Description", "Page No."],
        data=tables_data,
        col_widths=[1.3, 4.2, 0.9]
    )


def _build_abbreviations_and_symbols(doc):
    """Renders the List of Abbreviations and Symbols."""
    add_front_matter_heading(doc, "LIST OF ABBREVIATIONS AND SYMBOLS")

    abbr_data = [
        ["RTRFQM", "Real-Time Retail Footfall and Queue Monitoring Framework"],
        ["YOLO", "You Only Look Once (Deep Learning Object Detection Architecture)"],
        ["MOT", "Multi-Object Tracking"],
        ["MOTA", "Multi-Object Tracking Accuracy"],
        ["MOTP", "Multi-Object Tracking Precision"],
        ["IDF1", "Identification F1-Score (Ratio of Correctly Identified Detections)"],
        ["FPS", "Frames Per Second (Inference & Processing Throughput)"],
        ["ROI", "Region of Interest (Defined Polygonal Surveillance Zone)"],
        ["DFD", "Data Flow Diagram (Process and Data Modeling Artifact)"],
        ["POS", "Point of Sale (Retail Transaction Processing Register)"],
        ["CCTV", "Closed-Circuit Television (Overhead Video Surveillance)"],
        ["RTSP", "Real-Time Streaming Protocol (IP Video Camera Stream)"],
        ["MAE", "Mean Absolute Error"],
        ["RMSE", "Root Mean Squared Error"],
        ["MAPE", "Mean Absolute Percentage Error"],
        ["ARIMA", "Autoregressive Integrated Moving Average (Time-Series Baseline)"],
        ["IoU", "Intersection over Union (Bounding Box Overlap Metric)"],
        ["CIoU", "Complete Intersection over Union (Spatial Bounding Box Loss)"],
        ["DFL", "Distribution Focal Loss (BBox Coordinate Regression Loss)"],
        ["BCE", "Binary Cross-Entropy (Class Classification Loss)"],
        ["DDL", "Data Definition Language (SQL Schema Definition)"],
        ["KPI", "Key Performance Indicator (Operational Metric)"],
        ["λ (Lambda)", "Mean Customer Arrival Rate (Customers per Minute)"],
        ["μ (Mu)", "Mean Cashier Service Rate (Customers per Minute per Counter)"],
        ["Lq", "Mean Checkout Queue Length (Waiting Customers excluding Service)"],
        ["Wq", "Mean Customer Waiting Time in Queue (Minutes)"],
        ["ρ (Rho)", "Checkout Server Utilization Factor (Arrival Rate / Service Capacity)"],
        ["c", "Number of Active Cashier Checkout Counters / Lanes"],
    ]

    add_styled_table(
        doc,
        headers=["Acronym / Symbol", "Full Form / Mathematical Definition"],
        data=abbr_data,
        col_widths=[2.0, 4.4]
    )
